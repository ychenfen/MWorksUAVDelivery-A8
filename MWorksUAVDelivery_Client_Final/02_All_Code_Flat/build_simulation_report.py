from __future__ import annotations

import json
import math
from pathlib import Path

import numpy as np
import pandas as pd
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
OUT = ROOT / "simulation_outputs"
FIG = OUT / "figures"
REPORT = ROOT / "四旋翼无人机增强控制仿真报告.docx"

BASE_FONT = "Times New Roman"
CJK_FONT = "SimSun"
HEADING_CJK_FONT = "SimHei"
COLOR_BLUE = RGBColor(0, 0, 0)
COLOR_DARK_BLUE = RGBColor(0, 0, 0)
COLOR_MUTED = RGBColor(0, 0, 0)
COLOR_INK = RGBColor(0, 0, 0)
COLOR_BORDER = "000000"
COLOR_HEADER_FILL = "FFFFFF"
COLOR_CALLOUT = "FFFFFF"
CONTENT_DXA = 9360


def load_font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont | ImageFont.ImageFont:
    candidates = [
        Path(r"C:\Windows\Fonts\simhei.ttf"),
        Path(r"C:\Windows\Fonts\simsun.ttc"),
        Path(r"C:\Windows\Fonts\msyhbd.ttc") if bold else Path(r"C:\Windows\Fonts\msyh.ttc"),
        Path(r"C:\Windows\Fonts\arial.ttf"),
    ]
    for path in candidates:
        if path.exists():
            return ImageFont.truetype(str(path), size=size)
    return ImageFont.load_default()


FONT_TITLE = load_font(34, bold=True)
FONT_SUBTITLE = load_font(24, bold=True)
FONT_LABEL = load_font(20)
FONT_SMALL = load_font(17)
FONT_TINY = load_font(15)


def read_csv(name: str) -> pd.DataFrame:
    return pd.read_csv(OUT / f"{name}.csv")


def read_metrics() -> dict:
    return json.loads((OUT / "metrics_summary.json").read_text(encoding="utf-8"))


def nice_bounds(values: np.ndarray, pad_ratio: float = 0.08) -> tuple[float, float]:
    low = float(np.nanmin(values))
    high = float(np.nanmax(values))
    if not math.isfinite(low) or not math.isfinite(high):
        return 0.0, 1.0
    if abs(high - low) < 1e-9:
        return low - 1.0, high + 1.0
    pad = (high - low) * pad_ratio
    return low - pad, high + pad


def downsample_points(x: np.ndarray, y: np.ndarray, max_points: int = 1600) -> tuple[np.ndarray, np.ndarray]:
    if len(x) <= max_points:
        return x, y
    step = max(1, int(math.ceil(len(x) / max_points)))
    return x[::step], y[::step]


def chart_frame(title: str, subtitle: str | None = None, size: tuple[int, int] = (1400, 900)):
    img = Image.new("RGB", size, "white")
    draw = ImageDraw.Draw(img)
    draw.text((48, 30), title, fill=(0, 0, 0), font=FONT_TITLE)
    if subtitle:
        draw.text((50, 76), subtitle, fill=(0, 0, 0), font=FONT_SMALL)
    return img, draw


def draw_axes(
    draw: ImageDraw.ImageDraw,
    plot: tuple[int, int, int, int],
    x_min: float,
    x_max: float,
    y_min: float,
    y_max: float,
    x_label: str,
    y_label: str,
    equal: bool = False,
):
    left, top, right, bottom = plot
    draw.rectangle(plot, outline=(210, 216, 224), width=2)

    for i in range(6):
        x = left + (right - left) * i / 5
        y = bottom - (bottom - top) * i / 5
        draw.line((x, top, x, bottom), fill=(235, 238, 242), width=1)
        draw.line((left, y, right, y), fill=(235, 238, 242), width=1)
        x_val = x_min + (x_max - x_min) * i / 5
        y_val = y_min + (y_max - y_min) * i / 5
        draw.text((x - 32, bottom + 10), f"{x_val:.1f}", fill=(0, 0, 0), font=FONT_TINY)
        draw.text((left - 72, y - 10), f"{y_val:.1f}", fill=(0, 0, 0), font=FONT_TINY)

    draw.text(((left + right) // 2 - 35, bottom + 42), x_label, fill=(0, 0, 0), font=FONT_SMALL)
    draw.text((left - 78, top - 30), y_label, fill=(0, 0, 0), font=FONT_SMALL)
    if equal:
        draw.text((right - 140, bottom + 42), "等比例", fill=(0, 0, 0), font=FONT_TINY)


def draw_category_axes(
    draw: ImageDraw.ImageDraw,
    plot: tuple[int, int, int, int],
    y_min: float,
    y_max: float,
    x_label: str,
    y_label: str,
):
    left, top, right, bottom = plot
    draw.rectangle(plot, outline=(210, 216, 224), width=2)
    for i in range(6):
        y = bottom - (bottom - top) * i / 5
        draw.line((left, y, right, y), fill=(235, 238, 242), width=1)
        y_val = y_min + (y_max - y_min) * i / 5
        draw.text((left - 72, y - 10), f"{y_val:.1f}", fill=(0, 0, 0), font=FONT_TINY)
    draw.text(((left + right) // 2 - 35, bottom + 82), x_label, fill=(0, 0, 0), font=FONT_SMALL)
    draw.text((left - 78, top - 30), y_label, fill=(0, 0, 0), font=FONT_SMALL)


def to_screen(
    x: np.ndarray,
    y: np.ndarray,
    plot: tuple[int, int, int, int],
    x_min: float,
    x_max: float,
    y_min: float,
    y_max: float,
) -> list[tuple[int, int]]:
    left, top, right, bottom = plot
    sx = left + (x - x_min) / (x_max - x_min) * (right - left)
    sy = bottom - (y - y_min) / (y_max - y_min) * (bottom - top)
    return list(zip(np.round(sx).astype(int), np.round(sy).astype(int)))


def draw_polyline(
    draw: ImageDraw.ImageDraw,
    x: np.ndarray,
    y: np.ndarray,
    plot: tuple[int, int, int, int],
    x_min: float,
    x_max: float,
    y_min: float,
    y_max: float,
    color: tuple[int, int, int],
    width: int = 4,
):
    x, y = downsample_points(np.asarray(x), np.asarray(y))
    pts = to_screen(x, y, plot, x_min, x_max, y_min, y_max)
    if len(pts) > 1:
        draw.line(pts, fill=color, width=width, joint="curve")


def draw_legend(draw: ImageDraw.ImageDraw, items: list[tuple[str, tuple[int, int, int]]], x: int, y: int):
    for idx, (label, color) in enumerate(items):
        yy = y + idx * 30
        draw.line((x, yy + 10, x + 44, yy + 10), fill=color, width=5)
        draw.text((x + 55, yy), label, fill=(0, 0, 0), font=FONT_SMALL)


def make_time_chart(
    title: str,
    subtitle: str,
    base: pd.DataFrame,
    enhanced: pd.DataFrame,
    y_cols: tuple[str, str, str],
    y_label: str,
    out_path: Path,
):
    ref_col, base_col, enh_col = y_cols
    img, draw = chart_frame(title, subtitle)
    plot = (115, 135, 1040, 790)
    values = np.concatenate([base[ref_col].to_numpy(), base[base_col].to_numpy(), enhanced[enh_col].to_numpy()])
    x_min, x_max = nice_bounds(base["time"].to_numpy(), 0.0)
    y_min, y_max = nice_bounds(values)
    draw_axes(draw, plot, x_min, x_max, y_min, y_max, "时间 / s", y_label)
    draw_polyline(draw, base["time"], base[ref_col], plot, x_min, x_max, y_min, y_max, (35, 87, 137), 4)
    draw_polyline(draw, base["time"], base[base_col], plot, x_min, x_max, y_min, y_max, (194, 82, 60), 3)
    draw_polyline(draw, enhanced["time"], enhanced[enh_col], plot, x_min, x_max, y_min, y_max, (56, 139, 89), 3)
    draw_legend(draw, [("参考轨迹", (35, 87, 137)), ("基线响应", (194, 82, 60)), ("增强响应", (56, 139, 89))], 1090, 180)
    img.save(out_path, quality=95)


def make_xy_chart(title: str, subtitle: str, base: pd.DataFrame, enhanced: pd.DataFrame, out_path: Path):
    img, draw = chart_frame(title, subtitle)
    plot = (120, 135, 1050, 790)
    x_values = np.concatenate([base["ref_x"].to_numpy(), base["pos_x"].to_numpy(), enhanced["pos_x"].to_numpy()])
    y_values = np.concatenate([base["ref_y"].to_numpy(), base["pos_y"].to_numpy(), enhanced["pos_y"].to_numpy()])
    x_min, x_max = nice_bounds(x_values)
    y_min, y_max = nice_bounds(y_values)

    # Preserve the path shape by enforcing one common scale on x and y.
    x_mid = (x_min + x_max) / 2
    y_mid = (y_min + y_max) / 2
    plot_w = plot[2] - plot[0]
    plot_h = plot[3] - plot[1]
    data_w = max(x_max - x_min, 1e-6)
    data_h = max(y_max - y_min, 1e-6)
    scale = max(data_w / plot_w, data_h / plot_h)
    x_span = scale * plot_w
    y_span = scale * plot_h
    x_min, x_max = x_mid - x_span / 2, x_mid + x_span / 2
    y_min, y_max = y_mid - y_span / 2, y_mid + y_span / 2

    draw_axes(draw, plot, x_min, x_max, y_min, y_max, "X / m", "Y / m", equal=True)
    draw_polyline(draw, base["ref_x"], base["ref_y"], plot, x_min, x_max, y_min, y_max, (35, 87, 137), 5)
    draw_polyline(draw, base["pos_x"], base["pos_y"], plot, x_min, x_max, y_min, y_max, (194, 82, 60), 3)
    draw_polyline(draw, enhanced["pos_x"], enhanced["pos_y"], plot, x_min, x_max, y_min, y_max, (56, 139, 89), 3)
    draw_legend(draw, [("参考轨迹", (35, 87, 137)), ("基线轨迹", (194, 82, 60)), ("增强轨迹", (56, 139, 89))], 1090, 180)
    img.save(out_path, quality=95)


def make_bar_chart(metrics: dict, improvements: dict, out_path: Path):
    img, draw = chart_frame("综合误差指标对比", "路径 RMSE：数值越低越好")
    plot = (130, 145, 1180, 735)
    scenarios = [("climb", "爬升"), ("spiral", "螺旋"), ("eight", "8 字")]
    base_vals = [metrics[f"{key}_baseline"]["path_rmse_m"] for key, _ in scenarios]
    enh_vals = [metrics[f"{key}_enhanced"]["path_rmse_m"] for key, _ in scenarios]
    y_min, y_max = 0, max(base_vals + enh_vals) * 1.22
    draw_category_axes(draw, plot, y_min, y_max, "场景", "RMSE / m")
    group_w = (plot[2] - plot[0]) / len(scenarios)
    bar_w = 105
    colors = [(194, 82, 60), (56, 139, 89)]
    for idx, (key, label) in enumerate(scenarios):
        center = plot[0] + group_w * (idx + 0.5)
        for j, val in enumerate([base_vals[idx], enh_vals[idx]]):
            x0 = center + (-bar_w - 10 if j == 0 else 10)
            x1 = x0 + bar_w
            y0 = plot[3] - (val - y_min) / (y_max - y_min) * (plot[3] - plot[1])
            draw.rectangle((x0, y0, x1, plot[3]), fill=colors[j])
            draw.text((x0 + 10, y0 - 28), f"{val:.3f}", fill=(0, 0, 0), font=FONT_TINY)
        draw.text((center - 26, plot[3] + 34), label, fill=(0, 0, 0), font=FONT_SMALL)
        imp = improvements[key]["path_rmse_m_reduction_pct"]
        draw.text((center - 56, plot[1] + 20), f"降低 {imp:.1f}%", fill=(35, 105, 74), font=FONT_SMALL)
    draw_legend(draw, [("基线", colors[0]), ("增强", colors[1])], 1215, 210)
    img.save(out_path, quality=95)


def make_attitude_chart(metrics: dict, out_path: Path):
    img, draw = chart_frame("姿态安全边界检查", "最大 roll/pitch 与 15 deg 约束线")
    plot = (130, 145, 1180, 735)
    keys = ["climb_baseline", "climb_enhanced", "spiral_baseline", "spiral_enhanced", "eight_baseline", "eight_enhanced"]
    labels = ["爬升基线", "爬升增强", "螺旋基线", "螺旋增强", "8字基线", "8字增强"]
    values = [max(metrics[k]["max_roll_deg"], metrics[k]["max_pitch_deg"]) for k in keys]
    y_min, y_max = 0, max(max(values), 15) * 1.25
    draw_category_axes(draw, plot, y_min, y_max, "模型", "角度 / deg")
    limit_y = plot[3] - (15 - y_min) / (y_max - y_min) * (plot[3] - plot[1])
    draw.line((plot[0], limit_y, plot[2], limit_y), fill=(150, 40, 40), width=3)
    draw.text((plot[2] - 105, limit_y - 28), "15 deg", fill=(150, 40, 40), font=FONT_TINY)
    group_w = (plot[2] - plot[0]) / len(keys)
    bar_w = 85
    for idx, (label, val, key) in enumerate(zip(labels, values, keys)):
        center = plot[0] + group_w * (idx + 0.5)
        x0 = center - bar_w / 2
        y0 = plot[3] - (val - y_min) / (y_max - y_min) * (plot[3] - plot[1])
        color = (56, 139, 89) if "enhanced" in key else (194, 82, 60)
        draw.rectangle((x0, y0, x0 + bar_w, plot[3]), fill=color)
        draw.text((x0 + 6, y0 - 28), f"{val:.1f}", fill=(0, 0, 0), font=FONT_TINY)
        draw.text((center - 34, plot[3] + 34), label[:2], fill=(0, 0, 0), font=FONT_TINY)
        draw.text((center - 26, plot[3] + 56), label[2:], fill=(0, 0, 0), font=FONT_TINY)
    img.save(out_path, quality=95)


def set_run_font(run, name: str = BASE_FONT, east_asia: str = CJK_FONT, size: float | None = None,
                 color: RGBColor | None = None, bold: bool | None = None, italic: bool | None = None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), east_asia)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_style_font(style, size: float, color: RGBColor | None = None, bold: bool | None = None,
                   before: float = 0, after: float = 6, line_spacing: float = 1.10,
                   east_asia: str = CJK_FONT):
    style.font.name = BASE_FONT
    style._element.rPr.rFonts.set(qn("w:ascii"), BASE_FONT)
    style._element.rPr.rFonts.set(qn("w:hAnsi"), BASE_FONT)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), east_asia)
    style.font.size = Pt(size)
    if color:
        style.font.color.rgb = color
    if bold is not None:
        style.font.bold = bold
    pf = style.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line_spacing


def paragraph_border_bottom(paragraph, color="000000", size="12", space="1"):
    p = paragraph._p
    p_pr = p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), space)
    bottom.set(qn("w:color"), color)
    p_bdr.append(bottom)


def shade_cell(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(table, top=80, start=120, bottom=80, end=120):
    tbl_pr = table._tbl.tblPr
    tbl_cell_mar = tbl_pr.find(qn("w:tblCellMar"))
    if tbl_cell_mar is None:
        tbl_cell_mar = OxmlElement("w:tblCellMar")
        tbl_pr.append(tbl_cell_mar)
    for m, v in [("top", top), ("start", start), ("bottom", bottom), ("end", end)]:
        node = tbl_cell_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tbl_cell_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa: list[int], indent_dxa: int = 120):
    table.autofit = False
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")

    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    old_grid = tbl.find(qn("w:tblGrid"))
    if old_grid is not None:
        tbl.remove(old_grid)
    grid = OxmlElement("w:tblGrid")
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    tbl.insert(0, grid)

    for row in table.rows:
        for idx, width in enumerate(widths_dxa):
            if idx >= len(row.cells):
                continue
            cell = row.cells[idx]
            cell.width = Inches(width / 1440)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    set_cell_margins(table)


def set_table_borders(table, color=COLOR_BORDER):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ["top", "left", "bottom", "right", "insideH", "insideV"]:
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def add_text(paragraph, text: str, size: float | None = None, bold: bool | None = None,
             color: RGBColor | None = None, italic: bool | None = None,
             east_asia: str = CJK_FONT):
    run = paragraph.add_run(text)
    set_run_font(run, size=size, bold=bold, color=color, italic=italic, east_asia=east_asia)
    return run


def add_para(doc: Document, text: str = "", style: str | None = None, after: float | None = None,
             before: float | None = None):
    p = doc.add_paragraph(style=style)
    if text:
        add_text(p, text)
    if after is not None:
        p.paragraph_format.space_after = Pt(after)
    if before is not None:
        p.paragraph_format.space_before = Pt(before)
    return p


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_text(paragraph, "第 ", size=9, color=COLOR_MUTED)
    run = paragraph.add_run()
    fld_char_1 = OxmlElement("w:fldChar")
    fld_char_1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char_2 = OxmlElement("w:fldChar")
    fld_char_2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char_1)
    run._r.append(instr_text)
    run._r.append(fld_char_2)
    set_run_font(run, size=9, color=COLOR_MUTED)
    add_text(paragraph, " 页", size=9, color=COLOR_MUTED)


def add_metric_table(doc: Document, metrics: dict):
    rows = [
        ("爬升", "climb_baseline", "climb_enhanced"),
        ("螺旋爬升", "spiral_baseline", "spiral_enhanced"),
        ("8 字航迹", "eight_baseline", "eight_enhanced"),
    ]
    headers = ["场景", "模型", "RMSE / m", "平均误差 / m", "最大误差 / m", "最大姿态 / deg", "15 deg 超限点", "样本数"]
    table = doc.add_table(rows=1, cols=len(headers))
    set_table_geometry(table, [1350, 1500, 1120, 1300, 1300, 1400, 1320, 1250])
    set_table_borders(table)
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        shade_cell(cell, COLOR_HEADER_FILL)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_text(p, header, size=9.5, bold=True, color=COLOR_INK)
    for scenario_label, base_key, enh_key in rows:
        for model_label, key in [("基线", base_key), ("增强", enh_key)]:
            m = metrics[key]
            vals = [
                scenario_label,
                model_label,
                f"{m['path_rmse_m']:.4f}",
                f"{m['mean_path_error_m']:.4f}",
                f"{m['max_path_error_m']:.4f}",
                f"{m['max_attitude_norm_deg']:.2f}",
                str(m["angle_limit_15deg_violations"]),
                str(m["samples"]),
            ]
            cells = table.add_row().cells
            for i, val in enumerate(vals):
                p = cells[i].paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i != 0 else WD_ALIGN_PARAGRAPH.LEFT
                add_text(p, val, size=9.2, color=COLOR_INK)
    return table


def add_improvement_table(doc: Document, improvements: dict):
    rows = [
        ("爬升", improvements["climb"]),
        ("螺旋爬升", improvements["spiral"]),
        ("8 字航迹", improvements["eight"]),
    ]
    headers = ["场景", "RMSE 降低", "平均误差降低", "最大误差降低", "超调降低", "控制能量变化", "姿态范数变化"]
    table = doc.add_table(rows=1, cols=len(headers))
    set_table_geometry(table, [1450, 1250, 1400, 1400, 1150, 1350, 1360])
    set_table_borders(table)
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        shade_cell(cell, COLOR_HEADER_FILL)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_text(p, header, size=9.2, bold=True, color=COLOR_INK)
    for label, imp in rows:
        vals = [
            label,
            f"{imp['path_rmse_m_reduction_pct']:.2f}%",
            f"{imp['mean_path_error_m_reduction_pct']:.2f}%",
            f"{imp['max_path_error_m_reduction_pct']:.2f}%",
            f"{imp.get('z_overshoot_pct_reduction_pct', float('nan')):.2f}%" if "z_overshoot_pct_reduction_pct" in imp else "-",
            f"{imp['control_energy_proxy_reduction_pct']:.2f}%",
            f"{imp['max_attitude_norm_deg_reduction_pct']:.2f}%",
        ]
        cells = table.add_row().cells
        for i, val in enumerate(vals):
            p = cells[i].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_text(p, val, size=9.1, color=COLOR_INK)
    return table


def add_parameter_table(doc: Document):
    rows = [
        ("PID1", "KP=6.1, KI=0, KD=0.12", "高度/外环响应增强"),
        ("PID3 / PID4", "KP=2.0, KI=0, KD=1.35", "滚转/俯仰通道阻尼增强"),
        ("PID5 / PID6", "KP=17.0, KI=0, KD=2.05", "位置跟踪误差收敛增强"),
        ("PID7", "KP=10.0, KI=5.5, KD=5.0", "航向/综合通道稳态修正"),
        ("limiter1 / limiter2", "uMax=9.5/57.3, uMin=-9.5/57.3", "姿态命令限幅"),
        ("limiter3 / limiter4 / limiter5", "uMax=6.1, uMin=-6.1", "控制量限幅"),
    ]
    headers = ["对象", "增强版参数", "目的"]
    table = doc.add_table(rows=1, cols=3)
    set_table_geometry(table, [2100, 3650, 3610])
    set_table_borders(table)
    for i, header in enumerate(headers):
        shade_cell(table.rows[0].cells[i], COLOR_HEADER_FILL)
        p = table.rows[0].cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_text(p, header, size=9.5, bold=True, color=COLOR_INK)
    for row in rows:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            p = cells[i].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            add_text(p, val, size=9.2, color=COLOR_INK)


def add_caption(doc: Document, text: str):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(8)
    add_text(p, text, size=9.2, color=COLOR_MUTED)


def add_figure(doc: Document, image_path: Path, caption: str, width: float = 6.2):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(str(image_path), width=Inches(width))
    add_caption(doc, caption)


def setup_document() -> Document:
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    set_style_font(doc.styles["Normal"], 11, RGBColor(0, 0, 0), before=0, after=6, line_spacing=1.10)
    set_style_font(doc.styles["Title"], 22, RGBColor(0, 0, 0), bold=True, before=0, after=8, line_spacing=1.10, east_asia=HEADING_CJK_FONT)
    set_style_font(doc.styles["Subtitle"], 12, RGBColor(0, 0, 0), before=0, after=14, line_spacing=1.10)
    set_style_font(doc.styles["Heading 1"], 15, RGBColor(0, 0, 0), bold=True, before=14, after=7, line_spacing=1.10, east_asia=HEADING_CJK_FONT)
    set_style_font(doc.styles["Heading 2"], 13, RGBColor(0, 0, 0), bold=True, before=10, after=5, line_spacing=1.10, east_asia=HEADING_CJK_FONT)
    set_style_font(doc.styles["Heading 3"], 12, RGBColor(0, 0, 0), bold=True, before=8, after=4, line_spacing=1.10, east_asia=HEADING_CJK_FONT)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    add_text(header, "四旋翼无人机增强控制仿真报告", size=9, color=COLOR_MUTED)
    footer = section.footer.paragraphs[0]
    add_page_number(footer)
    return doc


def build_report(metrics_data: dict, figure_paths: dict[str, Path]):
    metrics = metrics_data["metrics"]
    improvements = metrics_data["improvements"]
    doc = setup_document()

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_text(title, "四旋翼无人机增强控制仿真报告", size=22, bold=True, color=RGBColor(0, 0, 0), east_asia=HEADING_CJK_FONT)
    subtitle = doc.add_paragraph(style="Subtitle")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_text(subtitle, "基于 MoHub QuadrotorModel 与 MWorks/Sysplorer 的航迹跟踪验证", size=12, color=RGBColor(0, 0, 0))

    metadata = [
        ("基础模型", "MoHub model/2355，QuadrotorModel，本地保存为 QuadrotorModel/package.mo"),
        ("仿真内容", "爬升、螺旋爬升、8 字航迹三类任务的基线/增强控制对比"),
        ("仿真平台", "MWorks/Sysplorer，Modelica 方程模型，Dassl 求解器"),
        ("数据来源", "simulation_outputs/*.csv 与 simulation_outputs/metrics_summary.json"),
        ("报告日期", "2026-05-30"),
    ]
    for label, value in metadata:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        add_text(p, f"{label}：", size=10.5, bold=True, color=COLOR_INK)
        add_text(p, value, size=10.5, color=COLOR_INK)
    rule = doc.add_paragraph()
    paragraph_border_bottom(rule, color="000000", size="8")

    doc.add_heading("1. 结论摘要", level=1)
    add_para(
        doc,
        "增强控制版本在三类任务中均降低路径跟踪 RMSE：爬升降低 15.39%，螺旋爬升降低 1.01%，8 字航迹降低 11.26%。其中爬升和 8 字任务收益最稳定；螺旋任务平均误差降低 12.79%，但最大姿态和约束超限点仍需后续安全控制器进一步处理。",
    )
    p = add_para(doc)
    add_text(
        p,
        "本次增强未改写 MoHub 基础包，而是在 QuadrotorTask.mo 中通过继承示例模型并覆盖 PID 与限幅参数形成对照组。该方式保持基础模型不变，便于复现实验和对比分析。",
    )

    doc.add_heading("2. 模型来源与任务分解", level=1)
    add_para(
        doc,
        "基础模型来自 MoHub 的 QuadrotorModel（项目 ID 2355，仓库 ID 1822）。公开接口可取得 Modelica 源码，本次已保存到本地 QuadrotorModel/package.mo；二进制贴图/可视化资源下载接口需要登录授权，因此本报告的仿真以 Modelica 方程模型和导出变量为依据。",
    )
    add_para(
        doc,
        "根据任务要求，仿真结构划分为五部分：四旋翼物理对象、参考轨迹、基线与增强控制器、安全限幅与约束检查、误差和稳定性评价。本次可执行版本将安全约束先落实为限幅参数和姿态超限统计，未单独实现 CBF/ERG/PPC 求解器。"
    )

    doc.add_heading("3. 仿真模型与增强方案", level=1)
    add_para(
        doc,
        "QuadrotorTask.mo 定义了六个可仿真模型：ClimbBaseline/ClimbEnhanced、SpiralBaseline/SpiralEnhanced、EightBaseline/EightEnhanced。基线模型分别继承 QuadrotorModel.Examples.Example1、Example2、Example3；增强模型在相同任务上覆盖控制器参数和限幅参数。",
    )
    doc.add_heading("增强参数", level=2)
    add_parameter_table(doc)

    doc.add_heading("4. 仿真设置", level=1)
    add_para(
        doc,
        "爬升和螺旋爬升仿真时间为 0-50 s，8 字航迹仿真时间为 0-120 s；采样间隔 0.01 s，容差 1e-4，求解算法为 Dassl。导出变量包含参考位置 ref_x/ref_y/ref_z、实际位置 pos_x/pos_y/pos_z、姿态 roll/pitch/yaw 以及控制输出 u1-u4。",
    )
    add_para(
        doc,
        "评价指标包括三轴 RMSE、路径 RMSE、平均路径误差、最大路径误差、爬升超调、2% 稳定时间、最大姿态范数、15 deg 姿态约束超限点数，以及控制输出能量代理指标。",
    )

    doc.add_heading("5. 仿真结果", level=1)
    doc.add_heading("主要指标", level=2)
    add_metric_table(doc, metrics)
    add_caption(doc, "表 1  基线与增强模型主要误差及安全指标")
    add_improvement_table(doc, improvements)
    add_caption(doc, "表 2  增强模型相对基线的改善比例；负值表示该项相对增加")

    add_figure(doc, figure_paths["rmse"], "图 1  三类任务路径 RMSE 对比", 6.2)
    add_figure(doc, figure_paths["climb_z"], "图 2  爬升任务高度跟踪响应", 6.2)
    add_figure(doc, figure_paths["spiral_xy"], "图 3  螺旋爬升任务 XY 航迹跟踪", 6.2)
    add_figure(doc, figure_paths["eight_xy"], "图 4  8 字任务 XY 航迹跟踪", 6.2)
    add_figure(doc, figure_paths["attitude"], "图 5  最大姿态与 15 deg 约束线对比", 6.2)

    doc.add_heading("6. 结果分析", level=1)
    add_para(
        doc,
        f"爬升任务中，路径 RMSE 从 {metrics['climb_baseline']['path_rmse_m']:.4f} m 降至 {metrics['climb_enhanced']['path_rmse_m']:.4f} m，平均路径误差降低 {improvements['climb']['mean_path_error_m_reduction_pct']:.2f}%，最大误差降低 {improvements['climb']['max_path_error_m_reduction_pct']:.2f}%。高度超调由 {metrics['climb_baseline']['z_overshoot_pct']:.3f}% 降至 {metrics['climb_enhanced']['z_overshoot_pct']:.3f}%，2% 稳定时间均为 {metrics['climb_baseline']['z_settling_time_2pct_s']:.1f} s。",
    )
    add_para(
        doc,
        f"螺旋爬升任务中，增强版路径 RMSE 小幅降低 {improvements['spiral']['path_rmse_m_reduction_pct']:.2f}%，平均误差降低 {improvements['spiral']['mean_path_error_m_reduction_pct']:.2f}%。最大误差几乎不变，说明该任务的瞬态偏差主要受参考轨迹初始段和模型动态限制影响，单纯参数调节难以完全消除。",
    )
    add_para(
        doc,
        f"8 字任务中，增强版路径 RMSE 从 {metrics['eight_baseline']['path_rmse_m']:.4f} m 降至 {metrics['eight_enhanced']['path_rmse_m']:.4f} m，平均误差降低 {improvements['eight']['mean_path_error_m_reduction_pct']:.2f}%，最大误差降低 {improvements['eight']['max_path_error_m_reduction_pct']:.2f}%。该场景没有出现 15 deg 姿态约束超限，说明增强参数对周期性水平航迹较适配。",
    )

    doc.add_heading("7. 安全性与局限", level=1)
    add_para(
        doc,
        f"15 deg 姿态约束统计显示：爬升增强版出现 {metrics['climb_enhanced']['angle_limit_15deg_violations']} 个超限点，螺旋增强版出现 {metrics['spiral_enhanced']['angle_limit_15deg_violations']} 个超限点，8 字增强版为 {metrics['eight_enhanced']['angle_limit_15deg_violations']} 个。增强参数提升了跟踪精度，但部分场景以更激进的瞬态姿态为代价，因此若任务要求严格安全边界，应继续加入显式 CBF/ERG/PPC 安全层，而不是只依赖 PID 和限幅器。",
    )
    add_para(
        doc,
        "控制输出 u1-u4 的超阈值计数在各模型中都较高，本报告将其作为控制量幅值代理指标，而不直接等同于真实电机饱和或硬件故障。若需要工程级安全结论，应进一步确认基础模型中 controller3_2.y/y1/y2/y3 与实际旋翼推力或转速之间的物理映射。",
    )

    doc.add_heading("8. 结论与后续工作", level=1)
    add_para(doc, "本次仿真表明，参数增强方案能够改善四旋翼模型在爬升和 8 字航迹中的跟踪精度，对螺旋爬升任务的平均误差也有改善。增强方案的代价是部分场景姿态峰值增大，因此后续工作应重点完善显式安全约束控制，并进一步确认控制输出与实际旋翼执行量之间的对应关系。")

    doc.save(REPORT)


def main():
    FIG.mkdir(parents=True, exist_ok=True)
    metrics_data = read_metrics()
    climb_b = read_csv("climb_baseline")
    climb_e = read_csv("climb_enhanced")
    spiral_b = read_csv("spiral_baseline")
    spiral_e = read_csv("spiral_enhanced")
    eight_b = read_csv("eight_baseline")
    eight_e = read_csv("eight_enhanced")

    figure_paths = {
        "climb_z": FIG / "climb_z_tracking.png",
        "spiral_xy": FIG / "spiral_xy_tracking.png",
        "eight_xy": FIG / "eight_xy_tracking.png",
        "rmse": FIG / "rmse_comparison.png",
        "attitude": FIG / "attitude_safety.png",
    }
    make_time_chart("爬升任务高度跟踪", "参考高度、基线响应与增强响应", climb_b, climb_e, ("ref_z", "pos_z", "pos_z"), "Z / m", figure_paths["climb_z"])
    make_xy_chart("螺旋爬升 XY 航迹", "参考轨迹、基线轨迹与增强轨迹", spiral_b, spiral_e, figure_paths["spiral_xy"])
    make_xy_chart("8 字 XY 航迹", "参考轨迹、基线轨迹与增强轨迹", eight_b, eight_e, figure_paths["eight_xy"])
    make_bar_chart(metrics_data["metrics"], metrics_data["improvements"], figure_paths["rmse"])
    make_attitude_chart(metrics_data["metrics"], figure_paths["attitude"])
    build_report(metrics_data, figure_paths)
    print(REPORT)


if __name__ == "__main__":
    main()
